#!/usr/bin/env python
#-*- coding: utf-8 -*-
#
from docx import *
import os

def B1():
    pass
if __name__=="__main__":
    try:
        os.makedirs("./ZScore")
    except:
        pass

    document = Document()
    table = document.add_table(10, 2 ,style="Table Grid")

    table.cell(0,0).text = '学号'
    table.cell(1,0).text = '姓名'
    table.cell(2,0).text = '思想品德素质得分明细'
    table.cell(3,0).text = '思想品德得分（10%）'
    table.cell(4,0).text = '身心素质得分明细'
    table.cell(5,0).text = '身心素质得分（5%）'
    table.cell(6,0).text = '创新实践能力得分明细'
    table.cell(7,0).text = '创新实践能力得分（10%）'
    table.cell(8,0).text = '学院特色得分明细'
    table.cell(9,0).text = '学院特色得分（5%）'


    document.save("./学号_姓名综测成绩导入模板.docx")
