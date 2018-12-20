#!/usr/bin/env python
#-*- coding: utf-8 -*-
#
#B2.py 综测成绩汇总
#1.遍历ZScore文件夹中的"xxxx学号_xxxx姓名综测成绩导入模板.docx"
#2.生成"综测成绩汇总.xls"
#   格式:学号 姓名 思想品德素质得分明细 思想品德得分（10%） 身心素质得分明细 身心素质得分（5%） 创新实践能力得分明细 创新实践能力得分（10%） 学院特色得分明细 学院特色得分（5%）
#
from docx import *
import os
import xlwt

def B2():
    pass
if __name__=="__main__":
    path = "./ZScore"
    files = os.listdir(path)
    workbook = xlwt.Workbook()
    sheet1 = workbook.add_sheet('sheet1',cell_overwrite_ok=True)

    style = xlwt.easyxf('align: wrap on,horz center,vert center')
    sheet1.col(0).width = 3300
    sheet1.col(2).width = 6000
    sheet1.col(4).width = 6000
    sheet1.col(6).width = 6000
    sheet1.col(8).width = 6000


    flag=0
    a=1
    for file in files:
        f = os.path.basename(file)
        Path=os.path.join(path,f)
        try:
            document = Document(Path)
            tables = document.tables 
            table = tables[0]

            if flag==0:
                for i in range(len(table.rows)):
                    sheet1.write(0,i,table.cell(i,0).text,style)
                    print(table.cell(i,0).text)
                flag=1
            
            for i in range(len(table.rows)):
                sheet1.write(a,i,table.cell(i,1).text,style)
                print(table.cell(i,1).text)
            a=a+1   
        except:
            pass
        
    workbook.save('./综测成绩汇总.xls')
